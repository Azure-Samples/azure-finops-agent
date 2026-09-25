import base64
import importlib.util
import io
from pathlib import Path
import unittest
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "src/Dashboard/AI/Tools/Resources/report_export.py"
SPEC = importlib.util.spec_from_file_location("report_export", SCRIPT)
REPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(REPORT)


class ReportTests(unittest.TestCase):
    def request(self, format_name):
        return {"format": format_name, "data": {"title": "Synthetic report", "source": "Synthetic test data", "sheets": [{"name": "Costs", "columns": ["name", "cost"], "rows": [["=1+1", 42.73], ["<script>bad()</script>", 10]], "sourceRowCount": 2}]}}

    def test_xlsx_preserves_numbers_and_neutralizes_formulas(self):
        result = REPORT.render(self.request("xlsx"))
        workbook = load_workbook(io.BytesIO(base64.b64decode(result["contentBase64"])))
        self.assertEqual(workbook.active["A2"].data_type, "s")
        self.assertEqual(workbook.active["A2"].value, "'=1+1")
        self.assertEqual(workbook.active["B2"].value, 42.73)
        self.assertEqual(workbook.active.max_row, 3)

    def test_csv_has_safe_formula_cells(self):
        content = base64.b64decode(REPORT.render(self.request("csv"))["contentBase64"]).decode("utf-8-sig")
        self.assertIn("'=1+1", content)
        self.assertIn("42.73", content)

    def test_html_escapes_data_and_has_filter(self):
        content = base64.b64decode(REPORT.render(self.request("html"))["contentBase64"]).decode("utf-8")
        self.assertNotIn("<script>bad()</script>", content)
        self.assertIn("&lt;script&gt;", content)
        self.assertIn('id="filter"', content)

    def test_declared_coverage_must_match_delivered_rows(self):
        request = self.request("html")
        request["data"]["sheets"][0]["sourceRowCount"] = 113
        with self.assertRaises(ValueError):
            REPORT.render(request)


if __name__ == "__main__":
    unittest.main()