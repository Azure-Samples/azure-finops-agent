import base64
import csv
import html
import io
import json
import math
import sys


def safe_cell(value):
    if value is not None and not isinstance(value, (str, int, float, bool)):
        raise ValueError("Cells must be scalar values")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Numeric cells must be finite")
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


THEME = """
:root {color-scheme:light;--cp-bg:#f7f4ef;--cp-bg-elevated:#fcfbf8;--cp-surface:#ffffff;--cp-surface-soft:#f5f5f5;--cp-border:#dedede;--cp-border-strong:#919191;--cp-text:#242424;--cp-text-muted:#5c5c5c;--cp-text-soft:#6f6f6f;--cp-accent:#b11f4b;--cp-accent-hover:#9a1a41;--cp-accent-soft:rgba(177,31,75,.08);--cp-accent-fg:#ffffff;--cp-success:#16a34a;--cp-danger:#dc2626;--cp-warning:#f59e0b;--cp-link:#0078d4;--cp-shadow:0 18px 48px rgba(0,0,0,.12);--cp-overlay:rgba(255,255,255,.8);--cp-panel:rgba(255,255,255,.86);--cp-panel-strong:rgba(255,255,255,.96);--cp-sheen:rgba(255,255,255,.55);--cp-highlight:rgba(177,31,75,.12)}
html[data-theme="dark"] {color-scheme:dark;--cp-bg:#3d3b3a;--cp-bg-elevated:#343231;--cp-surface:#292929;--cp-surface-soft:#2e2e2e;--cp-border:#474747;--cp-border-strong:#5f5f5f;--cp-text:#dedede;--cp-text-muted:#919191;--cp-text-soft:#b0b0b0;--cp-accent:#fd8ea1;--cp-accent-hover:#fb7b91;--cp-accent-soft:rgba(253,142,161,.14);--cp-accent-fg:#1a1a1a;--cp-success:#4ade80;--cp-danger:#f87171;--cp-warning:#fbbf24;--cp-link:#4da6ff;--cp-shadow:0 18px 48px rgba(0,0,0,.32);--cp-overlay:rgba(41,41,41,.88);--cp-panel:rgba(41,41,41,.72);--cp-panel-strong:rgba(41,41,41,.96);--cp-sheen:rgba(255,255,255,.04);--cp-highlight:rgba(253,142,161,.12)}
"""


def render(request):
    data = request["data"]
    sheets = data["sheets"]
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 10:
        raise ValueError("Provide 1 to 10 sheets")
    total = 0
    for sheet in sheets:
        columns, rows = sheet["columns"], sheet["rows"]
        if not 1 <= len(columns) <= 50 or not all(isinstance(column, str) and column for column in columns):
            raise ValueError("Provide 1 to 50 named columns")
        if len(set(columns)) != len(columns):
            raise ValueError("Column names must be unique")
        if not isinstance(rows, list) or any(not isinstance(row, list) or len(row) != len(columns) for row in rows):
            raise ValueError("Every row must match its columns")
        total += len(rows)
        if total > 5000:
            raise ValueError("Reports support at most 5000 total rows")
        if int(sheet.get("sourceRowCount", len(rows))) != len(rows):
            raise ValueError("Delivered rows do not reconcile with the declared source row count")
        for row in rows:
            for value in row:
                safe_cell(value)
    format_name = request["format"]
    if format_name == "csv":
        if len(sheets) != 1:
            raise ValueError("CSV supports one sheet; choose XLSX for a workbook")
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow([safe_cell(value) for value in sheets[0]["columns"]])
        writer.writerows([[safe_cell(value) for value in row] for row in sheets[0]["rows"]])
        content = buffer.getvalue().encode("utf-8-sig")
    elif format_name == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, sheet in enumerate(sheets):
            name = str(sheet.get("name") or f"Sheet{index + 1}")
            name = "".join(character for character in name if character not in "[]:*?/\\")[:31] or f"Sheet{index + 1}"
            worksheet = workbook.create_sheet(name)
            worksheet.append([safe_cell(value) for value in sheet["columns"]])
            for row in sheet["rows"]:
                worksheet.append([safe_cell(value) for value in row])
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="404040")
            for column in worksheet.columns:
                worksheet.column_dimensions[column[0].column_letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column[:100]) + 2))
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
    elif format_name == "html":
        escape = lambda value: html.escape("" if value is None else str(value), quote=True)
        sections = []
        for sheet in sheets:
            headings = "".join(f'<th scope="col">{escape(column)}</th>' for column in sheet["columns"])
            rows = "".join("<tr>" + "".join(f"<td>{escape(value)}</td>" for value in row) + "</tr>" for row in sheet["rows"])
            sections.append(f'<section><h2>{escape(sheet.get("name", "Data"))}</h2><p class="count">{len(sheet["rows"])} rows</p><div class="table-scroll"><table><thead><tr>{headings}</tr></thead><tbody>{rows}</tbody></table></div></section>')
        content = ('''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<script>(()=>{const param=new URLSearchParams(window.location.search).get("scoutTheme");const theme=param||(window.matchMedia("(prefers-color-scheme: dark)").matches?"dark":"light");document.documentElement.setAttribute("data-theme",theme)})();</script>
<title>''' + escape(data.get("title", "FinOps Report")) + "</title><style>" + THEME + '''
*{box-sizing:border-box}body{margin:0;padding:24px;font:14px "Segoe UI",Aptos,Calibri,sans-serif;letter-spacing:0;background:var(--cp-bg);color:var(--cp-text)}main{max-width:1440px;margin:auto}h1{font-size:28px;overflow-wrap:anywhere}h2{font-size:20px}p{color:var(--cp-text-muted);overflow-wrap:anywhere}input{width:min(100%,480px);padding:12px;border:1px solid var(--cp-border-strong);background:var(--cp-surface);color:var(--cp-text);border-radius:6px}section{margin-top:28px}.table-scroll{overflow:auto;border:1px solid var(--cp-border)}table{border-collapse:collapse;min-width:100%;background:var(--cp-surface)}th,td{text-align:left;padding:10px 12px;border-bottom:1px solid var(--cp-border);max-width:360px;overflow-wrap:anywhere}th{background:var(--cp-surface-soft);position:sticky;top:0}tr[hidden]{display:none}label{display:block;margin:20px 0 8px}@media print{input,label{display:none}body{padding:0}.table-scroll{overflow:visible}th{position:static}}
</style></head><body><main><h1>''' + escape(data.get("title", "FinOps Report")) + "</h1><p>" + escape(data.get("source", "Source timestamp unknown")) + '''</p><label for="filter">Filter rows</label><input type="search" id="filter" autocomplete="off">''' + "".join(sections) + '''</main><script>document.getElementById('filter').addEventListener('input',event=>{const query=event.target.value.toLowerCase();document.querySelectorAll('section').forEach(section=>{const rows=[...section.querySelectorAll('tbody tr')];rows.forEach(row=>row.hidden=!row.textContent.toLowerCase().includes(query));section.querySelector('.count').textContent=rows.filter(row=>!row.hidden).length+' / '+rows.length+' rows';});});</script></body></html>''').encode("utf-8")
    else:
        raise ValueError("Unsupported report format")
    return {"ok": True, "rows": total, "contentBase64": base64.b64encode(content).decode("ascii")}


if __name__ == "__main__":
    try:
        print(json.dumps(render(json.loads(sys.stdin.read()))))
    except (ValueError, TypeError, KeyError):
        print(json.dumps({"ok": False, "error": "Invalid report data, row coverage, or format"}))